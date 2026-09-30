"""
RAG 文档解析。

支持：
    txt / md / html / pdf / docx / xlsx

本模块只解析已保存到服务端的文件。
上传大小、文件类型和存储路径校验由 API 负责。

PDF 使用文本提取，不执行 OCR。
扫描版 PDF 如无文本层，需要单独的 OCR 流程。
"""

from dataclasses import dataclass
from pathlib import Path
import re

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass(frozen=True)
class ParsedSection:
    """解析后的一个文档片段。"""

    text: str
    source_page: int | None = None
    metadata: dict | None = None


SUPPORTED_SUFFIXES = frozenset({
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".xlsx",
})


MARKDOWN_HEADING_PATTERN = re.compile(
    r"^(#{1,6})\s+(.+?)\s*$"
)

QUESTION_HEADING_PATTERN = re.compile(
    r"(第\s*\d+\s*题|应该如何回答|怎么回答|如何回答|是什么|为什么|区别|怎么讲|怎么解释)"
)


def parse_text_file(
    path: Path,
) -> list[ParsedSection]:
    """解析 UTF-8 文本文件。"""

    text = path.read_text(
        encoding="utf-8-sig"
    )

    return [
        ParsedSection(text=text)
    ]


def parse_markdown_file(
    path: Path,
) -> list[ParsedSection]:
    """
    按 Markdown 问题标题拆分文档。

    题库类文档通常一个标题对应一道题，
    这样入库时能保留题目和答案的边界。
    """

    text = path.read_text(
        encoding="utf-8-sig"
    )

    sections = []
    current_lines = []
    current_title = None
    current_level = None

    def flush_current():
        nonlocal current_lines
        if not current_lines:
            return

        section_text = "\n".join(
            current_lines
        ).strip()
        if section_text:
            sections.append(
                ParsedSection(
                    text=section_text,
                    metadata={
                        "heading": current_title,
                        "heading_level": (
                            current_level
                        ),
                    }
                    if current_title
                    else None,
                )
            )
        current_lines = []

    for line in text.splitlines():
        heading = MARKDOWN_HEADING_PATTERN.match(line)
        if heading and QUESTION_HEADING_PATTERN.search(
            heading.group(2)
        ):
            flush_current()
            current_level = len(
                heading.group(1)
            )
            current_title = heading.group(2).strip()
            current_lines.append(line)
            continue

        current_lines.append(line)

    flush_current()

    return sections or [
        ParsedSection(text=text)
    ]


def parse_html_file(
    path: Path,
) -> list[ParsedSection]:
    """移除脚本和样式后提取 HTML 文本。"""

    html = path.read_text(
        encoding="utf-8-sig"
    )

    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    for element in soup([
        "script",
        "style",
        "noscript",
    ]):
        element.decompose()

    text = soup.get_text(
        separator="\n",
        strip=True,
    )

    return [
        ParsedSection(text=text)
    ]


def parse_pdf_file(
    path: Path,
) -> list[ParsedSection]:
    """按 PDF 页码提取文本。"""

    reader = PdfReader(
        str(path)
    )

    sections = []

    for page_index, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = (
            page.extract_text()
            or ""
        ).strip()

        if not text:
            continue

        sections.append(
            ParsedSection(
                text=text,
                source_page=page_index,
            )
        )

    return sections


def parse_docx_file(
    path: Path,
) -> list[ParsedSection]:
    """提取 Word 正文与表格文本。"""

    document = Document(
        str(path)
    )

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                parts.append(
                    " | ".join(cells)
                )

    return [
        ParsedSection(
            text="\n".join(parts)
        )
    ]


def parse_xlsx_file(
    path: Path,
) -> list[ParsedSection]:
    """按工作表提取 Excel 单元格文本。"""

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    sections = []

    try:
        for worksheet in workbook.worksheets:
            lines = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = [
                    ""
                    if value is None
                    else str(value)
                    for value in row
                ]

                if any(values):
                    lines.append(
                        " | ".join(values)
                    )

            if lines:
                sections.append(
                    ParsedSection(
                        text="\n".join(lines),
                        metadata={
                            "sheet": worksheet.title,
                        },
                    )
                )

    finally:
        workbook.close()

    return sections


def parse_document(
    path: str | Path,
) -> list[ParsedSection]:
    """
    根据扩展名选择解析器。

    不接受目录或不存在的路径。
    """

    source = Path(path)

    if (
        not source.is_file()
        or source.is_symlink()
    ):
        raise ValueError(
            "文档路径不存在或不是普通文件"
        )

    suffix = source.suffix.lower()

    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"不支持的文件类型：{suffix}"
        )

    if suffix == ".txt":
        return parse_text_file(source)

    if suffix in {
        ".md",
        ".markdown",
    }:
        return parse_markdown_file(source)

    if suffix in {
        ".html",
        ".htm",
    }:
        return parse_html_file(source)

    if suffix == ".pdf":
        return parse_pdf_file(source)

    if suffix == ".docx":
        return parse_docx_file(source)

    if suffix == ".xlsx":
        return parse_xlsx_file(source)

    raise ValueError(
        f"未找到解析器：{suffix}"
    )
